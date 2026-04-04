from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from pathlib import Path

import fitz
import httpx
import psycopg
from dotenv import load_dotenv
from pptx import Presentation
from qdrant_client import QdrantClient
from qdrant_client.http.models import FieldCondition, Filter, MatchValue

BACKEND_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = BACKEND_DIR.parent
ASSETS_DIR = BACKEND_DIR / "test_assets"
LOGS_DIR = BACKEND_DIR / "logs"
BASE_URL = "http://127.0.0.1:8000"

load_dotenv(BACKEND_DIR / ".env")

DATABASE_URL = os.environ["DATABASE_URL"]
PSYCOPG_DATABASE_URL = DATABASE_URL.replace("postgresql+psycopg://", "postgresql://", 1)
QDRANT_URL = os.environ["QDRANT_URL"]
QDRANT_COLLECTION_NAME = os.environ.get("QDRANT_COLLECTION_NAME", "document_chunks")


def assert_true(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def create_sample_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (72, 72),
        (
            "Quarterly Retention Report\n"
            "Churn decreased to 5 percent in Q2 2025.\n"
            "Proactive outreach improved retention.\n"
        ),
    )
    page = doc.new_page()
    page.insert_text(
        (72, 72),
        (
            "Recommendations\n"
            "The report recommends monthly health checks and customer education.\n"
        ),
    )
    doc.save(path)
    doc.close()


def create_sample_pptx(path: Path) -> None:
    presentation = Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Vector Database Lecture"
    slide.placeholders[1].text = (
        "Qdrant stores embeddings and supports payload filters for similarity search."
    )

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Lifecycle Notes"
    slide.placeholders[1].text = (
        "Delete vectors by document_id so stale chunks do not appear in retrieval results."
    )

    presentation.save(path)


def start_server() -> tuple[subprocess.Popen[str], Path]:
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    log_path = LOGS_DIR / "smoke_server.log"
    log_path.write_text("", encoding="utf-8")
    log_file = log_path.open("a", encoding="utf-8")
    process = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", "8000"],
        cwd=BACKEND_DIR,
        stdout=log_file,
        stderr=subprocess.STDOUT,
        text=True,
    )
    return process, log_path


def wait_for_health(process: subprocess.Popen[str], log_path: Path, timeout_seconds: int = 60) -> dict:
    deadline = time.time() + timeout_seconds
    with httpx.Client(timeout=5.0) as client:
        while time.time() < deadline:
            if process.poll() is not None:
                log_output = log_path.read_text(encoding="utf-8")
                raise RuntimeError(f"FastAPI exited early with code {process.returncode}.\n{log_output}")
            try:
                response = client.get(f"{BASE_URL}/api/health")
                response.raise_for_status()
                return response.json()
            except Exception:
                time.sleep(1)
    raise TimeoutError("FastAPI health endpoint did not become ready in time.")


def stop_server(process: subprocess.Popen[str]) -> None:
    if process.poll() is not None:
        return
    process.terminate()
    try:
        process.wait(timeout=10)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=10)


def db_fetch_one(query: str, params: tuple = ()) -> tuple | None:
    with psycopg.connect(PSYCOPG_DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(query, params)
            return cursor.fetchone()


def db_fetch_value(query: str, params: tuple = ()) -> int:
    row = db_fetch_one(query, params)
    return int(row[0]) if row else 0


def qdrant_count(client: QdrantClient, document_id: str) -> int:
    result = client.count(
        collection_name=QDRANT_COLLECTION_NAME,
        count_filter=Filter(
            must=[
                FieldCondition(
                    key="document_id",
                    match=MatchValue(value=document_id),
                )
            ]
        ),
        exact=True,
    )
    return int(result.count)


def upload_document(client: httpx.Client, path: Path, content_type: str) -> dict:
    with path.open("rb") as file_handle:
        response = client.post(
            f"{BASE_URL}/api/documents/upload",
            files={"file": (path.name, file_handle, content_type)},
        )
    response.raise_for_status()
    return response.json()


def main() -> int:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    pdf_path = ASSETS_DIR / "retention_report.pdf"
    pptx_path = ASSETS_DIR / "vector_notes.pptx"
    create_sample_pdf(pdf_path)
    create_sample_pptx(pptx_path)

    report: dict[str, object] = {}
    server_process, log_path = start_server()

    try:
        health = wait_for_health(server_process, log_path)
        report["health"] = health
        assert_true(health["status"] == "ok", "Health endpoint did not return status=ok.")
        assert_true(health["vector_store"] == "qdrant", "Vector store is not configured as qdrant.")

        qdrant = QdrantClient(url=QDRANT_URL, check_compatibility=False)
        assert_true(qdrant.collection_exists(QDRANT_COLLECTION_NAME), "Qdrant collection was not initialized.")

        with httpx.Client(timeout=20.0) as client:
            existing_documents = client.get(f"{BASE_URL}/api/documents")
            existing_documents.raise_for_status()
            for item in existing_documents.json()["items"]:
                delete_existing = client.delete(f"{BASE_URL}/api/documents/{item['id']}")
                assert_true(
                    delete_existing.status_code == 204,
                    f"Failed to clean up pre-existing document {item['id']}.",
                )

            pdf_upload = upload_document(client, pdf_path, "application/pdf")
            pptx_upload = upload_document(
                client,
                pptx_path,
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
            report["uploads"] = {
                "pdf": pdf_upload,
                "pptx": pptx_upload,
            }

            pdf_id = pdf_upload["document_id"]
            pptx_id = pptx_upload["document_id"]

            pdf_detail = client.get(f"{BASE_URL}/api/documents/{pdf_id}")
            pdf_detail.raise_for_status()
            pdf_detail_json = pdf_detail.json()

            pptx_detail = client.get(f"{BASE_URL}/api/documents/{pptx_id}")
            pptx_detail.raise_for_status()
            pptx_detail_json = pptx_detail.json()

            report["document_details"] = {
                "pdf": pdf_detail_json,
                "pptx": pptx_detail_json,
            }

            assert_true(pdf_detail_json["status"] == "indexed", "PDF status is not indexed.")
            assert_true(pptx_detail_json["status"] == "indexed", "PPTX status is not indexed.")
            assert_true(pdf_detail_json["page_count"] == 2, "PDF page_count should be 2.")
            assert_true(pptx_detail_json["slide_count"] == 2, "PPTX slide_count should be 2.")
            assert_true(Path(pdf_detail_json["storage_path"]).exists(), "Stored PDF file is missing.")
            assert_true(Path(pptx_detail_json["storage_path"]).exists(), "Stored PPTX file is missing.")

            pdf_chunk_count = db_fetch_value("SELECT COUNT(*) FROM chunks WHERE document_id = %s", (pdf_id,))
            pptx_chunk_count = db_fetch_value("SELECT COUNT(*) FROM chunks WHERE document_id = %s", (pptx_id,))
            assert_true(pdf_chunk_count > 0, "PDF did not create chunks.")
            assert_true(pptx_chunk_count > 0, "PPTX did not create chunks.")

            pdf_vector_count = qdrant_count(qdrant, pdf_id)
            pptx_vector_count = qdrant_count(qdrant, pptx_id)
            assert_true(pdf_vector_count == pdf_chunk_count, "PDF vectors were not fully inserted into Qdrant.")
            assert_true(pptx_vector_count == pptx_chunk_count, "PPTX vectors were not fully inserted into Qdrant.")

            document_list = client.get(f"{BASE_URL}/api/documents")
            document_list.raise_for_status()
            list_json = document_list.json()
            report["document_list"] = list_json
            filenames = {item["filename"] for item in list_json["items"]}
            assert_true("retention_report.pdf" in filenames, "PDF is missing from document list.")
            assert_true("vector_notes.pptx" in filenames, "PPTX is missing from document list.")

            pdf_qa = client.post(
                f"{BASE_URL}/api/chat/ask",
                json={
                    "question": "What churn percentage does the report mention?",
                    "mode": "private_only",
                    "selected_document_ids": [pdf_id],
                },
            )
            pdf_qa.raise_for_status()
            pdf_qa_json = pdf_qa.json()

            assert_true(pdf_qa_json["route_used"] == "private_only", "PDF QA route_used is incorrect.")
            assert_true(pdf_qa_json["retrieved_chunks"], "PDF QA did not retrieve any chunks.")
            assert_true(
                any(citation.get("page_number") == 1 for citation in pdf_qa_json["citations"]),
                "PDF citations are missing page metadata.",
            )
            assert_true(
                any(chunk["filename"] == "retention_report.pdf" for chunk in pdf_qa_json["retrieved_chunks"]),
                "PDF QA did not retrieve the expected document.",
            )
            assert_true(
                "5 percent" in pdf_qa_json["answer"].lower() or "5 percent" in json.dumps(pdf_qa_json).lower(),
                "PDF QA did not surface the expected churn value.",
            )

            pptx_qa = client.post(
                f"{BASE_URL}/api/chat/ask",
                json={
                    "question": "What do the slides say about deleting vectors by document_id?",
                    "mode": "private_only",
                    "file_type": "pptx",
                },
            )
            pptx_qa.raise_for_status()
            pptx_qa_json = pptx_qa.json()
            report["qa"] = {
                "pdf": pdf_qa_json,
                "pptx_filtered": pptx_qa_json,
            }

            assert_true(
                pptx_qa_json["route_used"] == "private_only",
                "PPTX filtered QA route_used is incorrect.",
            )
            assert_true(
                pptx_qa_json["retrieved_chunks"],
                "PPTX filtered QA did not retrieve any chunks.",
            )
            assert_true(
                all(chunk["filename"].endswith(".pptx") for chunk in pptx_qa_json["retrieved_chunks"]),
                "file_type filter did not restrict results to PPTX.",
            )
            assert_true(
                any(citation.get("slide_number") for citation in pptx_qa_json["citations"]),
                "PPTX citations are missing slide metadata.",
            )

            pdf_storage_path = Path(pdf_detail_json["storage_path"])
            delete_response = client.delete(f"{BASE_URL}/api/documents/{pdf_id}")
            assert_true(delete_response.status_code == 204, "Delete endpoint did not return 204.")

            remaining_document_rows = db_fetch_value("SELECT COUNT(*) FROM documents WHERE id = %s", (pdf_id,))
            remaining_chunk_rows = db_fetch_value("SELECT COUNT(*) FROM chunks WHERE document_id = %s", (pdf_id,))
            remaining_vectors = qdrant_count(qdrant, pdf_id)
            report["delete_checks"] = {
                "remaining_document_rows": remaining_document_rows,
                "remaining_chunk_rows": remaining_chunk_rows,
                "remaining_vectors": remaining_vectors,
                "storage_path_exists_after_delete": pdf_storage_path.exists(),
            }

            assert_true(remaining_document_rows == 0, "Deleted PDF still exists in SQL metadata.")
            assert_true(remaining_chunk_rows == 0, "Deleted PDF chunks still exist in SQL.")
            assert_true(remaining_vectors == 0, "Deleted PDF vectors still exist in Qdrant.")
            assert_true(not pdf_storage_path.exists(), "Deleted PDF file still exists on disk.")

            post_delete_qa = client.post(
                f"{BASE_URL}/api/chat/ask",
                json={
                    "question": "What churn percentage does the report mention?",
                    "mode": "private_only",
                },
            )
            post_delete_qa.raise_for_status()
            post_delete_qa_json = post_delete_qa.json()
            report["post_delete_qa"] = post_delete_qa_json

            assert_true(
                all(chunk["filename"] != "retention_report.pdf" for chunk in post_delete_qa_json["retrieved_chunks"]),
                "Deleted PDF still appears in retrieved chunks after deletion.",
            )
            assert_true(
                all(citation["title"] != "retention_report.pdf" for citation in post_delete_qa_json["citations"]),
                "Deleted PDF still appears in citations after deletion.",
            )

        print(json.dumps(report, indent=2))
        return 0
    finally:
        stop_server(server_process)


if __name__ == "__main__":
    raise SystemExit(main())
