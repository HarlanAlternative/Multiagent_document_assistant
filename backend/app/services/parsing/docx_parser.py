from pathlib import Path

from docx import Document as WordDocument

from app.services.parsing.base import ParsedDocument, ParsedUnit
from app.utils.text import normalize_whitespace


class DOCXParser:
    def parse(self, file_path: Path) -> ParsedDocument:
        document = WordDocument(file_path)
        units: list[ParsedUnit] = []
        current_section: str | None = None

        for paragraph in document.paragraphs:
            text = normalize_whitespace(paragraph.text)
            if not text:
                continue

            style_name = normalize_whitespace(getattr(paragraph.style, "name", ""))
            if style_name.lower().startswith("heading"):
                current_section = text
                continue

            units.append(
                ParsedUnit(
                    source_type="section",
                    text=text,
                    section_title=current_section,
                )
            )

        if not units:
            table_rows: list[str] = []
            for table in document.tables:
                for row in table.rows:
                    values = [normalize_whitespace(cell.text) for cell in row.cells]
                    values = [value for value in values if value]
                    if values:
                        table_rows.append(" | ".join(values))
            if table_rows:
                units.append(ParsedUnit(source_type="section", text="\n".join(table_rows)))

        return ParsedDocument(
            file_type="docx",
            source_path=file_path,
            units=units,
        )
