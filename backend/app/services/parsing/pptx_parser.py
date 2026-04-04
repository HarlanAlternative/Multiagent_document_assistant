from pathlib import Path

from pptx import Presentation

from app.services.parsing.base import ParsedDocument, ParsedUnit


class PPTXParser:
    def parse(self, file_path: Path) -> ParsedDocument:
        presentation = Presentation(file_path)
        units: list[ParsedUnit] = []

        for index, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text)
            combined_text = "\n".join(texts).strip()
            if not combined_text:
                continue
            units.append(
                ParsedUnit(
                    source_type="pptx_slide",
                    text=combined_text,
                    slide_number=index,
                )
            )

        return ParsedDocument(
            file_type="pptx",
            source_path=file_path,
            units=units,
            slide_count=len(presentation.slides),
        )
